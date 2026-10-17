"""Tạo bộ Excel mẫu nhỏ, cùng định dạng với 'TỔNG HỢP' thật, để kiểm tra tự động
mà không cần đưa dữ liệu thật lên GitHub.

Kết quả đúng được TÍNH TAY và ghi cứng vào TONG_HOP.xlsx mẫu (không lấy từ code đang kiểm tra):
  DA A (01/01–31/01/2026): THU 5.000.000 · CHI 1.000.000 + 1.200.000 − 100.000 (tiền thừa nhận về)
                            + D16 50.000 (thiếu ngày, vẫn tính) + chung 1.000.000 + 1.000.000 = 4.150.000
                            dòng THUẾ huỷ (SL 0) = 0
  DA B (từ 15/01/2026):    THU 3.000.000 · CHI   500.000 + chung 1.000.000 + 3.000.000            = 4.500.000
  Chung: 10/01 1tr (chỉ A chạy) · 20/01 2tr (A, B chia đôi) · 10/02 3tr (chỉ B)
"""
import datetime as dt
import os
import sys

import openpyxl

D = dt.datetime
HDR = ["STT", "NGAY", "TUAN", "LOAI", "NOI_DUNG", "CONG_TAC", "CONG_DOAN", "MA_CAY_TAY", "MA_CAY", "TEN_CAY",
       "DVT", "SL", "DON_GIA", "THANH_TIEN", "NGUON_TIEN", "PHAN_BO", "NHA_CC", "DONG_GOC", "KIEM_TRA",
       "GHI_CHU", "QUY_CACH"]
CAY = [("0", "TỔNG"), ("0.02", "THI CÔNG"), ("0.02.01", "NHÂN CÔNG"), ("0.02.01.01", "Thợ"),
       ("0.02.02", "VẬT TƯ THÔ"), ("0.02.02.02", "Thép"), ("1", "THU"), ("1.01", "Hợp đồng")]
MAP = [("THEP", "ALL", "0.02.02.02")]


def file_con(path, ma_da, nhanh, khoi_cong, dong):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NHAP"
    ws.append([f"SỔ NHẬP – {ma_da} – {nhanh}"])
    ws.append(["MA_DU_AN", ma_da, "NHANH", nhanh, "KHOI_CONG", khoi_cong])
    ws.append(["Chữ xanh = ô nhập."])
    ws.append(HDR)
    for i, d in enumerate(dong, start=5):
        r = {h: None for h in HDR}
        r.update(d)
        r["STT"] = "=ROW()-4"
        r["THANH_TIEN"] = f"=L{i}*M{i}"
        ws.append([r[h] for h in HDR])
    c = wb.create_sheet("CAY")
    c.append(["MA", "CAP", "TEN", "MA_CHA", "LA", "TEN_HIEN", "GHI_CHU", "TRANG_THAI"])
    for ma, ten in CAY:
        c.append([ma, None, ten])
    m = wb.create_sheet("MAP")
    m.append(["CONG_TAC", "CONG_DOAN", "MA_CAY", "TEN_CAY", "KEY"])
    for x in MAP:
        m.append(list(x))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def main(goc):
    th = os.path.join(goc, "TỔNG HỢP")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DU_AN"
    ws.append(["MA_DU_AN", "NAM", "TEN", "DIA_CHI", "KHOI_CONG", "KET_THUC", "THU_MUC", "GHI_CHU"])
    ws.append(["MAUA", 2026, "Dự án mẫu A", "1 Đường A", D(2026, 1, 1), D(2026, 1, 31), None, None])
    ws.append(["MAUB", 2026, "Dự án mẫu B", "2 Đường B", D(2026, 1, 15), None, None, None])
    os.makedirs(th, exist_ok=True)
    wb.save(os.path.join(th, "DANH_MUC_DU_AN.xlsx"))

    p = os.path.join(th, "2026")
    file_con(f"{p}/MAUA/0.02_THI_CONG/MAUA_0.02.02.02_THEP.xlsx", "MAUA", "0.02.02.02", D(2026, 1, 1), [
        dict(NGAY=D(2026, 1, 5), LOAI="CHI", NOI_DUNG="D10", CONG_TAC="THEP", CONG_DOAN="THEP", SL=10, DON_GIA=100000),
        dict(NGAY=D(2026, 1, 6), LOAI="CHI", NOI_DUNG="D12", MA_CAY_TAY="0.02.02.02", SL="=2*5", DON_GIA="=(120)*1000"),
        dict(LOAI="CHI", NOI_DUNG="ĐINH", MA_CAY_TAY="0.02.02.02"),  # thiếu cả SL, ĐG → đã huỷ (0 đ)
        dict(NGAY=D(2026, 1, 7), LOAI="CHI", NOI_DUNG="THUẾ HUỶ", MA_CAY_TAY="0.02.02.02", SL=0, DON_GIA=50000000),
        dict(NGAY=D(2026, 1, 8), LOAI="CHI", NOI_DUNG="TIỀN THỪA NHẬN VỀ", MA_CAY_TAY="0.02.02.02", SL=1, DON_GIA=-100000),
        dict(LOAI="CHI", NOI_DUNG="D16", MA_CAY_TAY="0.02.02.02", SL=1, DON_GIA=50000),  # có tiền, thiếu ngày → xem lại
    ])
    file_con(f"{p}/MAUA/1_THU/MAUA_1_THU.xlsx", "MAUA", "1", D(2026, 1, 1), [
        dict(NGAY=D(2026, 1, 2), LOAI="THU", NOI_DUNG="Hợp đồng", MA_CAY_TAY="1.01", SL=1, DON_GIA=5000000),
    ])
    file_con(f"{p}/MAUB/0.02_THI_CONG/MAUB_0.02.02.02_THEP.xlsx", "MAUB", "0.02.02.02", D(2026, 1, 15), [
        dict(NGAY=D(2026, 1, 20), LOAI="CHI", NOI_DUNG="D10", CONG_TAC="THEP", CONG_DOAN="ALL", SL=5, DON_GIA=100000),
    ])
    file_con(f"{p}/MAUB/1_THU/MAUB_1_THU.xlsx", "MAUB", "1", D(2026, 1, 15), [
        dict(NGAY=D(2026, 1, 16), LOAI="THU", NOI_DUNG="Hợp đồng", MA_CAY_TAY="1.01", SL=1, DON_GIA=3000000),
    ])
    file_con(f"{p}/APG_CHUNG/APG_CHUNG_2026_0.02.01.01_THO.xlsx", "CHUNG", "0.02.01.01", None, [
        dict(NGAY=D(2026, 1, 10), LOAI="CHI", NOI_DUNG="LƯƠNG T", MA_CAY_TAY="0.02.01.01", SL=1, DON_GIA=1000000, PHAN_BO="CHUNG"),
        dict(NGAY=D(2026, 1, 20), LOAI="CHI", NOI_DUNG="LƯƠNG T", MA_CAY_TAY="0.02.01.01", SL=1, DON_GIA=2000000, PHAN_BO="CHUNG"),
        dict(NGAY=D(2026, 2, 10), LOAI="CHI", NOI_DUNG="LƯƠNG T", MA_CAY_TAY="0.02.01.01", SL=1, DON_GIA=3000000, PHAN_BO="CHUNG"),
    ])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DU_AN"
    ws.append(["MA_DU_AN", "NAM", "TEN", "KHOI_CONG", "KET_THUC", "THU", "CHI"])
    ws.append(["MAUA", 2026, "Dự án mẫu A", D(2026, 1, 1), D(2026, 1, 31), 5000000, 4150000])
    ws.append(["MAUB", 2026, "Dự án mẫu B", D(2026, 1, 15), None, 3000000, 4500000])
    wb.save(os.path.join(th, "TONG_HOP.xlsx"))
    print("Đã tạo dữ liệu mẫu:", th)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "tests/mau")
