"""Đọc bộ file Excel 'TỔNG HỢP' của APG thành dữ liệu sạch để đưa vào PocketBase.

Không dùng giá trị đã lưu trong ô công thức (nhiều file chưa từng được Excel tính lại),
mà tự tính lại: tra MA_CAY qua MAP như công thức Excel, tính SL/ĐG từ công thức số học.
"""
from __future__ import annotations

import ast
import datetime as dt
import operator
import os
from collections import Counter, defaultdict
from dataclasses import dataclass, field

import openpyxl

# ---------------------------------------------------------------- công thức số học


_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul,
    ast.Div: operator.truediv, ast.USub: operator.neg, ast.UAdd: operator.pos,
}


def tinh_so(v):
    """Số, None, hoặc công thức số học kiểu '=(262.332)*1000*(1-0.15)'. Không chấp nhận gì khác."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise ValueError(f"không phải số: {v!r}")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith("="):
        s = s[1:]
    s = s.replace(",", ".")

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return float(n.value)
        if isinstance(n, ast.BinOp) and type(n.op) in _OPS:
            return _OPS[type(n.op)](ev(n.left), ev(n.right))
        if isinstance(n, ast.UnaryOp) and type(n.op) in _OPS:
            return _OPS[type(n.op)](ev(n.operand))
        raise ValueError(f"công thức không phải số học thuần: {v!r}")

    return ev(ast.parse(s, mode="eval"))


def ngay_iso(v):
    if v in (None, ""):
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"ngày không đọc được: {v!r}")


def chu(v):
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- mô hình


@dataclass
class CayMa:
    ma: str
    ten: str
    ghi_chu: str = ""

    @property
    def ma_cha(self):
        return "" if "." not in self.ma else self.ma.rsplit(".", 1)[0]

    @property
    def cap(self):
        return self.ma.count(".")


@dataclass
class DuAn:
    ma: str
    ten: str
    dia_chi: str
    khoi_cong: str
    ket_thuc: str
    ghi_chu: str


@dataclass
class GiaoDich:
    ma_import: str
    du_an: str          # mã dự án, '' nếu CHUNG
    loai: str
    ngay: str
    noi_dung: str
    ma_cay: str
    dvt: str
    so_luong: float | None
    don_gia: float | None
    nguon_tien: str
    phan_bo: str
    nha_cc: str
    quy_cach: str
    ghi_chu: str
    nguon_goc: str
    trang_thai: str = "HIEU_LUC"   # DA_HUY: giao dịch có thật nhưng đã huỷ, thành tiền = 0
    ly_do: list = field(default_factory=list)

    @property
    def thanh_tien(self):
        # giống hook trên server: làm tròn đến đồng; đã huỷ = 0
        if self.trang_thai == "DA_HUY":
            return 0
        return round((self.so_luong or 0) * (self.don_gia or 0))

    @property
    def thanh_tien_excel(self):
        return (self.so_luong or 0) * (self.don_gia or 0)


@dataclass
class BoDuLieu:
    cay: list
    map_cu: dict
    du_an: list
    giao_dich: list
    canh_bao: list


# ---------------------------------------------------------------- đọc


def _doc_cay_map(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    cay = [CayMa(chu(r[0]), chu(r[2]), chu(r[6]) if len(r) > 6 else "")
           for r in wb["CAY"].iter_rows(min_row=2, values_only=True) if r[0]]
    mp = {}
    for r in wb["MAP"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            mp[(chu(r[0]), chu(r[1]))] = chu(r[2])
    wb.close()
    return cay, mp


def _doc_danh_muc(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["DU_AN"]
    rows = ws.iter_rows(values_only=True)
    hdr = [chu(c) for c in next(rows)]
    ix = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for r in rows:
        ma = chu(r[ix["MA_DU_AN"]])
        if not ma:
            continue
        out.append(DuAn(ma, chu(r[ix["TEN"]]), chu(r[ix["DIA_CHI"]]),
                        ngay_iso(r[ix["KHOI_CONG"]]), ngay_iso(r[ix["KET_THUC"]]),
                        chu(r[ix["GHI_CHU"]])))
    wb.close()
    return out


def doc_thu_muc(goc: str) -> BoDuLieu:
    """goc = thư mục 'TỔNG HỢP' (chứa DANH_MUC_DU_AN.xlsx)."""
    canh_bao = []
    du_an = _doc_danh_muc(os.path.join(goc, "DANH_MUC_DU_AN.xlsx"))
    da = {d.ma: d for d in du_an}

    file_con = []
    for r, _, fs in os.walk(goc):
        for f in sorted(fs):
            if f.endswith(".xlsx") and not f.startswith("~$"):
                p = os.path.join(r, f)
                wb = openpyxl.load_workbook(p, read_only=True)
                if "NHAP" in wb.sheetnames:
                    file_con.append(p)
                wb.close()
    file_con.sort()
    if not file_con:
        raise SystemExit(f"Không thấy file con nào (sheet NHAP) trong {goc}")

    # cây + map phải giống nhau ở mọi file con
    cay, mp = _doc_cay_map(file_con[0])
    sig = tuple(c.ma for c in cay)
    for p in file_con[1:]:
        c2, m2 = _doc_cay_map(p)
        if tuple(c.ma for c in c2) != sig or m2 != mp:
            canh_bao.append(f"Cây mã/MAP ở {os.path.relpath(p, goc)} khác file đầu tiên – dùng bản của file đầu")
    co_con = Counter(c.ma_cha for c in cay if c.ma_cha)
    la = {c.ma: co_con[c.ma] == 0 for c in cay}

    gd = []
    for p in file_con:
        rel = os.path.relpath(p, goc)
        wb = openpyxl.load_workbook(p)  # cần công thức gốc, không dùng read_only
        ws = wb["NHAP"]
        r2 = [c.value for c in ws[2]]
        ma_da = chu(r2[r2.index("MA_DU_AN") + 1])
        nhanh = chu(r2[r2.index("NHANH") + 1])
        hdr = [chu(c.value) for c in ws[4]]
        ix = {h: i for i, h in enumerate(hdr) if h}
        for i, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            g = lambda k: row[ix[k]] if k in ix and ix[k] < len(row) else None
            nd = chu(g("NOI_DUNG"))
            if not nd:
                continue
            ma_tay = chu(g("MA_CAY_TAY"))
            ct, cd = chu(g("CONG_TAC")), chu(g("CONG_DOAN"))
            ma = ma_tay or mp.get((ct, cd)) or mp.get((ct, "ALL")) or (nhanh if la.get(nhanh) else "")
            loai = chu(g("LOAI")).upper()
            phan_bo = "CHUNG" if chu(g("PHAN_BO")).upper() == "CHUNG" or ma_da == "CHUNG" else "RIENG"
            ly_do = []
            try:
                sl = tinh_so(g("SL"))
                dg = tinh_so(g("DON_GIA"))
            except ValueError as ex:
                raise SystemExit(f"{rel} dòng {i}: {ex}")
            try:
                ngay = ngay_iso(g("NGAY"))
            except ValueError as ex:
                raise SystemExit(f"{rel} dòng {i}: {ex}")

            # lỗi cứng: không thể lưu dù có đánh dấu
            if not ma:
                raise SystemExit(f"{rel} dòng {i} ({nd}): chưa có mã cây")
            if ma not in la:
                raise SystemExit(f"{rel} dòng {i} ({nd}): mã {ma} không có trong cây")
            if not la[ma]:
                raise SystemExit(f"{rel} dòng {i} ({nd}): mã {ma} chưa xuống lá")
            if ma != nhanh and not ma.startswith(nhanh + "."):
                raise SystemExit(f"{rel} dòng {i} ({nd}): mã {ma} sai nhánh {nhanh}")
            if loai not in ("CHI", "THU"):
                raise SystemExit(f"{rel} dòng {i} ({nd}): LOAI = {loai!r}")
            if (loai == "THU") != (ma == "1" or ma.startswith("1.")):
                raise SystemExit(f"{rel} dòng {i} ({nd}): {loai} nhưng mã {ma}")
            if phan_bo == "RIENG" and ma_da not in da:
                raise SystemExit(f"{rel}: dự án {ma_da} không có trong DANH_MUC_DU_AN")
            # Quy ước đã chốt (18/10/2026):
            #  - thành tiền = 0 → giao dịch có thật nhưng đã huỷ: vẫn lưu, trạng thái DA_HUY
            #  - số âm → tiền thừa nhận về / trừ ứng / chiết khấu: hợp lệ, không đánh dấu
            da_huy = (sl or 0) * (dg or 0) == 0
            if phan_bo == "CHUNG" and not ngay and not da_huy:
                raise SystemExit(f"{rel} dòng {i} ({nd}): chi phí CHUNG thiếu ngày")

            # nghi vấn: vẫn chuyển, đánh dấu cần xem lại
            if not da_huy:
                if not ngay:
                    ly_do.append("thiếu ngày")
                if phan_bo == "RIENG" and ngay:
                    d = da[ma_da]
                    if ngay < d.khoi_cong:
                        ly_do.append(f"ngày {ngay} trước khởi công {d.khoi_cong}")
                    if d.ket_thuc and ngay > d.ket_thuc:
                        ly_do.append(f"ngày {ngay} sau kết thúc {d.ket_thuc}")

            gc = g("GHI_CHU")
            gd.append(GiaoDich(
                ma_import=f"{rel}#{i}",
                du_an="" if phan_bo == "CHUNG" else ma_da,
                loai=loai, ngay=ngay, noi_dung=nd, ma_cay=ma,
                dvt=chu(g("DVT")), so_luong=sl, don_gia=dg,
                nguon_tien=chu(g("NGUON_TIEN")).upper(), phan_bo=phan_bo,
                nha_cc=chu(g("NHA_CC")), quy_cach=chu(g("QUY_CACH")),
                ghi_chu=ngay_iso(gc) if isinstance(gc, (dt.date, dt.datetime)) else chu(gc),
                nguon_goc=chu(g("DONG_GOC")), ly_do=ly_do,
                trang_thai="DA_HUY" if da_huy else "HIEU_LUC",
            ))
        wb.close()

    # nghi trùng: cùng dự án, ngày, nội dung, SL, ĐG (chỉ xét giao dịch còn hiệu lực)
    nhom = defaultdict(list)
    for x in gd:
        if x.trang_thai == "HIEU_LUC":
            nhom[(x.du_an, x.ngay, x.noi_dung.upper(), x.so_luong, x.don_gia)].append(x)
    for xs in nhom.values():
        if len(xs) > 1:
            for x in xs:
                khac = ", ".join(f"dòng {y.ma_import.rsplit('#', 1)[1]}"
                                 + ("" if y.ma_import.rsplit('#', 1)[0] == x.ma_import.rsplit('#', 1)[0]
                                    else f" file {y.ma_import.rsplit('#', 1)[0]}")
                                 + (f" (sổ gốc {y.nguon_goc})" if y.nguon_goc else "")
                                 for y in xs if y is not x)
                x.ly_do.append(f"nghi trùng với {khac}")

    return BoDuLieu(cay=cay, map_cu=mp, du_an=du_an, giao_dich=gd, canh_bao=canh_bao)


def la_cua(cay):
    co_con = Counter(c.ma_cha for c in cay if c.ma_cha)
    return {c.ma: co_con[c.ma] == 0 for c in cay}
